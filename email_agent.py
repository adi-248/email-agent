import os
import base64
import re
import json
import requests

from datetime import datetime, timezone
from email.utils import parsedate_to_datetime
from html import unescape

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))

CREDENTIALS_FILE = os.path.join(
    PROJECT_DIR,
    "credentials",
    "credentials.json"
)

TOKEN_FILE = os.path.join(
    PROJECT_DIR,
    "credentials",
    "token.json"
)

REPORT_DIR = os.path.join(
    PROJECT_DIR,
    "reports"
)

OLLAMA_URL = "http://127.0.0.1:11434/api/generate"
OLLAMA_MODEL = "qwen2.5:3b"

# Start with 10. Once results are verified, increase gradually.
MAX_MESSAGES = 10

# Maximum body characters sent to AI for the current message.
MAX_BODY_FOR_AI = 6000

# Maximum characters of thread context sent to AI.
MAX_THREAD_CONTEXT_FOR_AI = 5000

# Confidence below this level will be flagged for review.
LOW_CONFIDENCE_THRESHOLD = 0.75


# ============================================================
# CONTROLLED AI VALUES
# ============================================================

ALLOWED_CATEGORIES = [
    "Business",
    "Finance",
    "Personal",
    "Marketing",
    "Newsletter",
    "OTP",
    "Bills",
    "Orders",
    "Banking",
    "Tax",
    "Government",
    "Legal",
    "Jobs",
    "Social Security",
    "Subscriptions",
    "Notifications",
    "Receipts",
    "Support",
    "Spam",
    "Other",
]

ALLOWED_PRIORITIES = [
    "Critical",
    "High",
    "Medium",
    "Low",
]

ALLOWED_ACTIONS = [
    "Reply",
    "Review",
    "Approve",
    "Pay",
    "Submit",
    "Call",
    "Attend",
    "Download",
    "Sign",
    "Upload",
    "Complete",
    "Renew",
    "Verify",
    "Book",
    "Cancel",
    "Track",
    "No action",
]

ALLOWED_YES_NO = [
    "Yes",
    "No",
]

ALLOWED_REPLY_VALUES = [
    "Yes",
    "No",
    "Maybe",
]

ALLOWED_RELEVANCE = [
    "High",
    "Medium",
    "Low",
]

ALLOWED_DEADLINE_TYPES = [
    "Exact",
    "Relative",
    "Inferred",
    "No deadline",
]

ALLOWED_THREAD_STATUS = [
    "Single message",
    "New conversation",
    "Existing conversation",
]


# ============================================================
# GMAIL PERMISSIONS
# ============================================================

# READ ONLY.
# This script cannot delete, move, archive, label or send email.
SCOPES = [
    "https://www.googleapis.com/auth/gmail.readonly"
]


# ============================================================
# FOLDER SETUP
# ============================================================

os.makedirs(REPORT_DIR, exist_ok=True)
os.makedirs(os.path.dirname(CREDENTIALS_FILE), exist_ok=True)


# ============================================================
# GMAIL AUTHENTICATION
# ============================================================

def authenticate_gmail():
    creds = None

    if os.path.exists(TOKEN_FILE):
        creds = Credentials.from_authorized_user_file(
            TOKEN_FILE,
            SCOPES
        )

    if not creds or not creds.valid:

        if creds and creds.expired and creds.refresh_token:
            print("Refreshing Gmail authentication...")
            creds.refresh(Request())

        else:
            print("Starting Gmail authentication...")
            print("A browser window will open.")

            if not os.path.exists(CREDENTIALS_FILE):
                raise FileNotFoundError(
                    f"\nGoogle OAuth credentials not found:\n"
                    f"{CREDENTIALS_FILE}\n\n"
                    f"Place credentials.json in the credentials folder."
                )

            flow = InstalledAppFlow.from_client_secrets_file(
                CREDENTIALS_FILE,
                SCOPES
            )

            creds = flow.run_local_server(port=0)

        with open(TOKEN_FILE, "w", encoding="utf-8") as token:
            token.write(creds.to_json())

    return build(
        "gmail",
        "v1",
        credentials=creds
    )


# ============================================================
# GMAIL MESSAGE LIST
# ============================================================

def get_message_ids(service):
    print("\nScanning Gmail mailbox...")

    message_ids = []
    page_token = None

    while True:
        response = service.users().messages().list(
            userId="me",
            maxResults=500,
            pageToken=page_token
        ).execute()

        for message in response.get("messages", []):
            message_ids.append(message["id"])

            if MAX_MESSAGES and len(message_ids) >= MAX_MESSAGES:
                return message_ids[:MAX_MESSAGES]

        page_token = response.get("nextPageToken")

        if not page_token:
            break

    return message_ids


# ============================================================
# HEADER EXTRACTION
# ============================================================

def get_header(headers, name):
    name = name.lower()

    for header in headers:
        if header.get("name", "").lower() == name:
            return header.get("value", "")

    return ""


# ============================================================
# EMAIL DATE
# ============================================================

def parse_email_date(date_string):
    if not date_string:
        return None

    try:
        dt = parsedate_to_datetime(date_string)

        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)

        return dt

    except Exception:
        return None


# ============================================================
# HTML CLEANING
# ============================================================

def clean_html(html):
    if not html:
        return ""

    html = re.sub(
        r"<script.*?>.*?</script>",
        " ",
        html,
        flags=re.IGNORECASE | re.DOTALL
    )

    html = re.sub(
        r"<style.*?>.*?</style>",
        " ",
        html,
        flags=re.IGNORECASE | re.DOTALL
    )

    html = re.sub(
        r"<br\s*/?>",
        "\n",
        html,
        flags=re.IGNORECASE
    )

    html = re.sub(
        r"</p>",
        "\n",
        html,
        flags=re.IGNORECASE
    )

    html = re.sub(r"<[^>]+>", " ", html)
    html = unescape(html)

    html = re.sub(r"\s+", " ", html).strip()

    return html


# ============================================================
# EMAIL BODY EXTRACTION
# ============================================================

def extract_text_from_payload(payload):
    plain_parts = []
    html_parts = []

    def process_part(part):
        mime_type = part.get("mimeType", "")
        body = part.get("body", {})
        data = body.get("data")

        if data:
            try:
                decoded = base64.urlsafe_b64decode(data).decode(
                    "utf-8",
                    errors="ignore"
                )

                if mime_type == "text/plain":
                    plain_parts.append(decoded)

                elif mime_type == "text/html":
                    html_parts.append(decoded)

            except Exception:
                pass

        for child in part.get("parts", []):
            process_part(child)

    process_part(payload)

    if plain_parts:
        text = "\n".join(plain_parts)

    elif html_parts:
        text = clean_html("\n".join(html_parts))

    else:
        text = ""

    return re.sub(r"\s+", " ", text).strip()


# ============================================================
# ATTACHMENT ANALYSIS
# ============================================================

def extract_attachment_info(payload):
    attachments = []

    def process_part(part):
        filename = part.get("filename", "")
        body = part.get("body", {})

        if filename:
            attachments.append({
                "filename": filename,
                "mime_type": part.get("mimeType", ""),
                "size": body.get("size", 0),
                "attachment_id": body.get("attachmentId", ""),
            })

        for child in part.get("parts", []):
            process_part(child)

    process_part(payload)

    return attachments


# ============================================================
# SIZE FORMATTING
# ============================================================

def format_size(bytes_value):
    if bytes_value == 0:
        return "0 B"

    units = ["B", "KB", "MB", "GB", "TB"]
    size = float(bytes_value)

    for unit in units:
        if size < 1024:
            return f"{size:.2f} {unit}"

        size /= 1024

    return f"{size:.2f} PB"


# ============================================================
# THREAD AWARENESS
# ============================================================

def get_thread_context(service, thread_id, current_message_id):
    """
    Returns:
      - thread message count
      - thread status
      - compact context from other messages in the same thread
    """

    if not thread_id:
        return {
            "thread_message_count": 1,
            "thread_status": "Single message",
            "thread_context": "",
        }

    try:
        thread = service.users().threads().get(
            userId="me",
            id=thread_id,
            format="full"
        ).execute()

        messages = thread.get("messages", [])
        count = len(messages)

        if count <= 1:
            status = "Single message"
        else:
            status = "Existing conversation"

        context_items = []

        for message in messages:
            if message.get("id") == current_message_id:
                continue

            payload = message.get("payload", {})
            headers = payload.get("headers", [])

            sender = get_header(headers, "From")
            subject = get_header(headers, "Subject")
            date_string = get_header(headers, "Date")
            body = extract_text_from_payload(payload)

            context_items.append(
                f"FROM: {sender}\n"
                f"DATE: {date_string}\n"
                f"SUBJECT: {subject}\n"
                f"BODY: {body[:1200]}"
            )

        thread_context = "\n\n--- THREAD MESSAGE ---\n\n".join(
            context_items
        )

        return {
            "thread_message_count": count,
            "thread_status": status,
            "thread_context": thread_context[:MAX_THREAD_CONTEXT_FOR_AI],
        }

    except Exception as error:
        return {
            "thread_message_count": 1,
            "thread_status": "Single message",
            "thread_context": f"Thread context unavailable: {error}",
        }


# ============================================================
# AI VALIDATION HELPERS
# ============================================================

def normalize_choice(value, allowed_values, fallback):
    if value is None:
        return fallback

    value_text = str(value).strip()

    for allowed in allowed_values:
        if value_text.lower() == allowed.lower():
            return allowed

    return fallback


def normalize_confidence(value):
    try:
        confidence = float(value)
    except Exception:
        return 0.0

    return max(0.0, min(confidence, 1.0))


def normalize_yes_no(value, fallback="No"):
    if isinstance(value, bool):
        return "Yes" if value else "No"

    return normalize_choice(
        value,
        ALLOWED_YES_NO,
        fallback
    )


def validate_ai_output(classification):
    """
    Validates AI output against our controlled vocabulary.
    Returns normalized data + validation warnings.
    """

    warnings = []

    category_raw = classification.get("category")
    category = normalize_choice(
        category_raw,
        ALLOWED_CATEGORIES,
        "Other"
    )

    if str(category_raw).strip().lower() != category.lower():
        warnings.append("Invalid category returned by AI")

    priority_raw = classification.get("priority")
    priority = normalize_choice(
        priority_raw,
        ALLOWED_PRIORITIES,
        "Low"
    )

    if str(priority_raw).strip().lower() != priority.lower():
        warnings.append("Invalid priority returned by AI")

    action_raw = classification.get("action")
    action = normalize_choice(
        action_raw,
        ALLOWED_ACTIONS,
        "No action"
    )

    if str(action_raw).strip().lower() != action.lower():
        warnings.append("Invalid action returned by AI")

    deadline_type_raw = classification.get("deadline_type")
    deadline_type = normalize_choice(
        deadline_type_raw,
        ALLOWED_DEADLINE_TYPES,
        "No deadline"
    )

    if str(deadline_type_raw).strip().lower() != deadline_type.lower():
        warnings.append("Invalid deadline type returned by AI")

    requires_reply_raw = classification.get("requires_reply")
    requires_reply = normalize_choice(
        requires_reply_raw,
        ALLOWED_REPLY_VALUES,
        "No"
    )

    if str(requires_reply_raw).strip().lower() != requires_reply.lower():
        warnings.append("Invalid requires_reply returned by AI")

    relevance_raw = classification.get("relevance")
    relevance = normalize_choice(
        relevance_raw,
        ALLOWED_RELEVANCE,
        "Low"
    )

    if str(relevance_raw).strip().lower() != relevance.lower():
        warnings.append("Invalid relevance returned by AI")

    is_automated = normalize_yes_no(
        classification.get("is_automated"),
        "No"
    )

    is_promotional = normalize_yes_no(
        classification.get("is_promotional"),
        "No"
    )

    is_urgent = normalize_yes_no(
        classification.get("is_urgent"),
        "No"
    )

    new_action_in_thread = normalize_yes_no(
        classification.get("new_action_in_thread"),
        "No"
    )

    confidence = normalize_confidence(
        classification.get("confidence", 0)
    )

    deadline = str(
        classification.get(
            "deadline",
            "No deadline"
        )
    ).strip() or "No deadline"

    deadline_text = str(
        classification.get(
            "deadline_text",
            "No deadline"
        )
    ).strip() or "No deadline"

    if deadline_type == "No deadline":
        deadline = "No deadline"
        deadline_text = "No deadline"

    action_description = str(
        classification.get(
            "action_description",
            "No action required"
        )
    ).strip() or "No action required"

    reason = str(
        classification.get(
            "reason",
            ""
        )
    ).strip()

    ai_status = "OK"

    if warnings or confidence < LOW_CONFIDENCE_THRESHOLD:
        ai_status = "REVIEW"

    return {
        "category": category,
        "priority": priority,
        "action": action,
        "action_description": action_description,
        "deadline": deadline,
        "deadline_type": deadline_type,
        "deadline_text": deadline_text,
        "requires_reply": requires_reply,
        "relevance": relevance,
        "is_automated": is_automated,
        "is_promotional": is_promotional,
        "is_urgent": is_urgent,
        "new_action_in_thread": new_action_in_thread,
        "reason": reason,
        "confidence": confidence,
        "ai_status": ai_status,
        "validation_warnings": "; ".join(warnings),
    }


# ============================================================
# QWEN CLASSIFICATION
# ============================================================

def classify_email_with_qwen(
    sender,
    recipient,
    subject,
    email_date,
    body,
    attachments,
    thread_context,
    thread_message_count,
    thread_status,
):
    attachment_text = ", ".join(
        item["filename"]
        for item in attachments
    ) or "None"

    current_time = datetime.now(timezone.utc)

    email_text = f"""
CURRENT UTC DATE/TIME:
{current_time.strftime("%Y-%m-%d %H:%M:%S UTC")}

EMAIL DATE:
{email_date.strftime("%Y-%m-%d %H:%M:%S %Z") if email_date else "Unknown"}

SENDER:
{sender}

RECIPIENT:
{recipient}

SUBJECT:
{subject}

ATTACHMENTS:
{attachment_text}

THREAD MESSAGE COUNT:
{thread_message_count}

THREAD STATUS:
{thread_status}

OTHER MESSAGES IN THIS THREAD:
{thread_context if thread_context else "None"}

CURRENT EMAIL BODY:
{body[:MAX_BODY_FOR_AI]}
"""

    prompt = f"""
You are an email intelligence agent.

Analyze the CURRENT EMAIL in the context of the thread.

Return ONLY valid JSON.
Do not use markdown.
Do not add explanations outside JSON.

Required JSON structure:

{{
  "category": "",
  "priority": "",
  "action": "",
  "action_description": "",
  "deadline": "",
  "deadline_type": "",
  "deadline_text": "",
  "requires_reply": "",
  "relevance": "",
  "is_automated": "",
  "is_promotional": "",
  "is_urgent": "",
  "new_action_in_thread": "",
  "reason": "",
  "confidence": 0.0
}}

CATEGORY must be exactly one of:

Business
Finance
Personal
Marketing
Newsletter
OTP
Bills
Orders
Banking
Tax
Government
Legal
Jobs
Social Security
Subscriptions
Notifications
Receipts
Support
Spam
Other

CATEGORY GUIDANCE:

Business:
Work, client, vendor, professional or commercial correspondence.

Finance:
General financial matters that are not specifically banking, tax, bills,
receipts or orders.

Personal:
Family, friends and genuinely personal communications.

Marketing:
Sales promotions, advertisements, offers, discounts or promotional mail.

Newsletter:
Recurring informational publications, digests or newsletters.

OTP:
One-time passwords, verification codes or login codes.

Bills:
Utility bills, invoices and payment-due notices.

Orders:
Order confirmations, shipping, delivery, returns or e-commerce activity.

Banking:
Bank account, credit/debit card, transaction, statement, loan or banking alerts.

Tax:
Tax notices, filing, tax payments, tax documents or tax-related communication.

Government:
Government departments, public authorities and official civic communication.

Legal:
Contracts, legal notices, legal correspondence or legal obligations.

Jobs:
Job applications, recruiters, interviews, job portals or employment opportunities.

Social Security:
Social security, pension, provident fund, retirement/social-benefit schemes
or similar official benefit communication.
Do NOT use this category for social media.

Subscriptions:
Subscription renewals, cancellations or subscription-service account notices.

Notifications:
General automated informational notifications that do not better fit another category.

Receipts:
Payment confirmations and receipts for completed purchases/payments.

Support:
Customer support, help desk, complaint or service-resolution communication.

Spam:
Clearly unwanted, deceptive, irrelevant or junk email.

Other:
Use only when no other category fits.

PRIORITY must be exactly one of:

Critical
High
Medium
Low

PRIORITY GUIDANCE:

Critical:
Immediate serious risk or consequence, such as security compromise,
fraud, legal emergency, account blocking, payment failure with major consequence,
or an action/deadline requiring immediate attention.

High:
Important action required soon, important business/financial matter,
payment, approval, reply or time-sensitive task.

Medium:
Relevant and useful, but not immediately urgent.

Low:
Routine FYI, marketing, newsletter, non-actionable notification,
or low-impact message.

ACTION must be exactly one of:

Reply
Review
Approve
Pay
Submit
Call
Attend
Download
Sign
Upload
Complete
Renew
Verify
Book
Cancel
Track
No action

ACTION_DESCRIPTION:
Briefly state exactly what the recipient should do.
If no action is required, use:
"No action required"

DEADLINE:
Use an exact ISO-style date/time when it can be reliably determined.
Examples:
2026-09-09
2026-09-09 17:00

If the message only says something relative such as "tomorrow",
"today", "within 3 days", interpret it using CURRENT UTC DATE/TIME
and EMAIL DATE where appropriate.

If there is no deadline, use:
"No deadline"

DEADLINE_TYPE must be exactly one of:

Exact
Relative
Inferred
No deadline

DEADLINE_TEXT:
Preserve the human wording that created the deadline, for example:
"tomorrow by 5 PM"
"within 3 days"
"September 15"
If none exists, use:
"No deadline"

REQUIRES_REPLY must be exactly one of:

Yes
No
Maybe

RELEVANCE must be exactly one of:

High
Medium
Low

IS_AUTOMATED must be:
Yes
No

IS_PROMOTIONAL must be:
Yes
No

IS_URGENT must be:
Yes
No

NEW_ACTION_IN_THREAD must be:
Yes
No

THREAD GUIDANCE:

Do not treat every message in a thread as a new task.

If earlier messages already requested an action and the current message
only says "thanks", "noted", "received", "okay", "done", or otherwise
does not create a new task, use:

"action": "No action"
"new_action_in_thread": "No"

If this current email adds or changes a required action,
use:
"new_action_in_thread": "Yes"

CONFIDENCE must be a number between 0 and 1.

REASON:
Briefly explain the classification, priority and action.

EMAIL TO ANALYZE:

{email_text}
"""

    try:
        response = requests.post(
            OLLAMA_URL,
            json={
                "model": OLLAMA_MODEL,
                "prompt": prompt,
                "stream": False,
                "format": "json",
            },
            timeout=180
        )

        response.raise_for_status()

        result = response.json()
        raw_response = result.get("response", "")

        if not raw_response:
            raise ValueError("Ollama returned an empty response.")

        classification = json.loads(raw_response)

        return validate_ai_output(classification)

    except requests.exceptions.ConnectionError:
        return ai_error_result(
            "Ollama is not running",
            "Cannot connect to Ollama."
        )

    except requests.exceptions.Timeout:
        return ai_error_result(
            "Qwen request timed out",
            "Ollama did not respond within the timeout."
        )

    except json.JSONDecodeError:
        return ai_error_result(
            "Invalid JSON from Qwen",
            "Qwen did not return valid JSON."
        )

    except Exception as error:
        return ai_error_result(
            "Classification failed",
            str(error)
        )


def ai_error_result(action_description, reason):
    return {
        "category": "Other",
        "priority": "Low",
        "action": "No action",
        "action_description": action_description,
        "deadline": "No deadline",
        "deadline_type": "No deadline",
        "deadline_text": "No deadline",
        "requires_reply": "No",
        "relevance": "Low",
        "is_automated": "No",
        "is_promotional": "No",
        "is_urgent": "No",
        "new_action_in_thread": "No",
        "reason": reason,
        "confidence": 0.0,
        "ai_status": "AI ERROR",
        "validation_warnings": reason,
    }


# ============================================================
# PROCESS ONE EMAIL
# ============================================================

def process_message(service, message_id):
    message = service.users().messages().get(
        userId="me",
        id=message_id,
        format="full"
    ).execute()

    payload = message.get("payload", {})
    headers = payload.get("headers", [])

    sender = get_header(headers, "From")
    recipient = get_header(headers, "To")
    subject = get_header(headers, "Subject")
    date_string = get_header(headers, "Date")

    email_date = parse_email_date(date_string)

    labels = message.get("labelIds", [])
    attachments = extract_attachment_info(payload)
    body = extract_text_from_payload(payload)

    attachment_count = len(attachments)

    attachment_size = sum(
        item["size"]
        for item in attachments
    )

    if email_date:
        now = datetime.now(timezone.utc)
        age_days = (now - email_date).days
    else:
        age_days = ""

    attachment_names = "; ".join(
        item["filename"]
        for item in attachments
    )

    attachment_types = "; ".join(
        item["mime_type"]
        for item in attachments
    )

    thread_id = message.get("threadId", "")

    thread_info = get_thread_context(
        service=service,
        thread_id=thread_id,
        current_message_id=message_id
    )

    print("\n" + "-" * 70)
    print(f"Subject : {subject}")
    print(f"Sender  : {sender}")
    print(
        f"Thread  : {thread_info['thread_message_count']} message(s)"
    )
    print("Sending to Qwen...")

    ai = classify_email_with_qwen(
        sender=sender,
        recipient=recipient,
        subject=subject,
        email_date=email_date,
        body=body,
        attachments=attachments,
        thread_context=thread_info["thread_context"],
        thread_message_count=thread_info["thread_message_count"],
        thread_status=thread_info["thread_status"],
    )

    print(f"Category   : {ai['category']}")
    print(f"Priority   : {ai['priority']}")
    print(f"Action     : {ai['action']}")
    print(f"Deadline   : {ai['deadline']}")
    print(f"Relevance  : {ai['relevance']}")
    print(f"Automated  : {ai['is_automated']}")
    print(f"Promotional: {ai['is_promotional']}")
    print(f"Urgent     : {ai['is_urgent']}")
    print(f"New Action : {ai['new_action_in_thread']}")
    print(f"Confidence : {ai['confidence']:.2f}")
    print(f"AI Status  : {ai['ai_status']}")

    return {
        "message_id": message_id,
        "thread_id": thread_id,
        "thread_message_count": thread_info["thread_message_count"],
        "thread_status": thread_info["thread_status"],

        "date": email_date.strftime(
            "%Y-%m-%d %H:%M:%S"
        ) if email_date else "",

        "age_days": age_days,
        "sender": sender,
        "recipient": recipient,
        "subject": subject,
        "labels": ", ".join(labels),

        "attachment_count": attachment_count,
        "attachment_names": attachment_names,
        "attachment_types": attachment_types,
        "attachment_size_bytes": attachment_size,
        "attachment_size": format_size(attachment_size),

        "body_preview": body[:2000],

        "category": ai["category"],
        "priority": ai["priority"],
        "action": ai["action"],
        "action_description": ai["action_description"],
        "deadline": ai["deadline"],
        "deadline_type": ai["deadline_type"],
        "deadline_text": ai["deadline_text"],
        "requires_reply": ai["requires_reply"],
        "relevance": ai["relevance"],
        "is_automated": ai["is_automated"],
        "is_promotional": ai["is_promotional"],
        "is_urgent": ai["is_urgent"],
        "new_action_in_thread": ai["new_action_in_thread"],
        "reason": ai["reason"],
        "confidence": ai["confidence"],
        "ai_status": ai["ai_status"],
        "validation_warnings": ai["validation_warnings"],
    }


# ============================================================
# EXCEL HELPERS
# ============================================================

def add_email_detail_headers(sheet):
    headers = [
        "Message ID",
        "Thread ID",
        "Thread Message Count",
        "Thread Status",
        "New Action In Thread",
        "Date",
        "Age (Days)",
        "Sender",
        "Recipient",
        "Subject",
        "Gmail Labels",
        "AI Category",
        "AI Priority",
        "AI Relevance",
        "AI Action",
        "Action Description",
        "AI Deadline",
        "Deadline Type",
        "Deadline Text",
        "Requires Reply",
        "Automated",
        "Promotional",
        "Urgent",
        "AI Reason",
        "AI Confidence",
        "AI Status",
        "Validation Warnings",
        "Attachment Count",
        "Attachment Names",
        "Attachment Types",
        "Attachment Size",
        "Body Preview",
    ]

    sheet.append(headers)


def append_email_detail_row(sheet, record):
    sheet.append([
        record["message_id"],
        record["thread_id"],
        record["thread_message_count"],
        record["thread_status"],
        record["new_action_in_thread"],
        record["date"],
        record["age_days"],
        record["sender"],
        record["recipient"],
        record["subject"],
        record["labels"],
        record["category"],
        record["priority"],
        record["relevance"],
        record["action"],
        record["action_description"],
        record["deadline"],
        record["deadline_type"],
        record["deadline_text"],
        record["requires_reply"],
        record["is_automated"],
        record["is_promotional"],
        record["is_urgent"],
        record["reason"],
        record["confidence"],
        record["ai_status"],
        record["validation_warnings"],
        record["attachment_count"],
        record["attachment_names"],
        record["attachment_types"],
        record["attachment_size"],
        record["body_preview"],
    ])


def populate_filtered_sheet(workbook, title, records):
    sheet = workbook.create_sheet(title)
    add_email_detail_headers(sheet)

    for record in records:
        append_email_detail_row(sheet, record)

    return sheet


# ============================================================
# EXCEL REPORT
# ============================================================

def create_excel_report(records):
    timestamp = datetime.now().strftime(
        "%Y-%m-%d_%H-%M-%S"
    )

    filename = f"Email_AI_V3_Report_{timestamp}.xlsx"

    filepath = os.path.join(
        REPORT_DIR,
        filename
    )

    workbook = Workbook()

    # ========================================================
    # SUMMARY
    # ========================================================

    summary = workbook.active
    summary.title = "Summary"

    total_emails = len(records)

    emails_with_attachments = sum(
        1
        for r in records
        if r["attachment_count"] > 0
    )

    total_attachments = sum(
        r["attachment_count"]
        for r in records
    )

    total_attachment_size = sum(
        r["attachment_size_bytes"]
        for r in records
    )

    action_required = sum(
        1
        for r in records
        if r["action"] != "No action"
    )

    high_priority = sum(
        1
        for r in records
        if r["priority"] in ["Critical", "High"]
    )

    low_confidence = sum(
        1
        for r in records
        if r["confidence"] < LOW_CONFIDENCE_THRESHOLD
        or r["ai_status"] != "OK"
    )

    category_counts = {}
    priority_counts = {}
    action_counts = {}
    relevance_counts = {}

    for record in records:
        category_counts[record["category"]] = (
            category_counts.get(record["category"], 0) + 1
        )

        priority_counts[record["priority"]] = (
            priority_counts.get(record["priority"], 0) + 1
        )

        action_counts[record["action"]] = (
            action_counts.get(record["action"], 0) + 1
        )

        relevance_counts[record["relevance"]] = (
            relevance_counts.get(record["relevance"], 0) + 1
        )

    summary.append(["SRNR EMAIL AGENT - V3 EMAIL INTELLIGENCE"])
    summary.append([])

    summary.append([
        "Report Generated",
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ])

    summary.append(["AI Model", OLLAMA_MODEL])
    summary.append(["Low Confidence Threshold", LOW_CONFIDENCE_THRESHOLD])
    summary.append([])

    summary.append(["Total Emails", total_emails])
    summary.append(["Action Required", action_required])
    summary.append(["Critical / High Priority", high_priority])
    summary.append(["Low Confidence / Review", low_confidence])
    summary.append(["Emails With Attachments", emails_with_attachments])
    summary.append(["Total Attachments", total_attachments])
    summary.append([
        "Total Attachment Size",
        format_size(total_attachment_size)
    ])

    summary.append([])
    summary.append(["CATEGORY", "COUNT"])

    for category in ALLOWED_CATEGORIES:
        summary.append([
            category,
            category_counts.get(category, 0)
        ])

    summary.append([])
    summary.append(["PRIORITY", "COUNT"])

    for priority in ALLOWED_PRIORITIES:
        summary.append([
            priority,
            priority_counts.get(priority, 0)
        ])

    summary.append([])
    summary.append(["ACTION", "COUNT"])

    for action in ALLOWED_ACTIONS:
        summary.append([
            action,
            action_counts.get(action, 0)
        ])

    summary.append([])
    summary.append(["RELEVANCE", "COUNT"])

    for relevance in ALLOWED_RELEVANCE:
        summary.append([
            relevance,
            relevance_counts.get(relevance, 0)
        ])

    # ========================================================
    # EMAIL DETAILS
    # ========================================================

    details = workbook.create_sheet("Email Details")
    add_email_detail_headers(details)

    for record in records:
        append_email_detail_row(details, record)

    # ========================================================
    # FILTERED WORKSHEETS
    # ========================================================

    action_records = [
        r for r in records
        if r["action"] != "No action"
    ]

    deadline_records = [
        r for r in records
        if r["deadline_type"] != "No deadline"
        and r["deadline"] != "No deadline"
    ]

    high_priority_records = [
        r for r in records
        if r["priority"] in ["Critical", "High"]
    ]

    low_confidence_records = [
        r for r in records
        if r["confidence"] < LOW_CONFIDENCE_THRESHOLD
        or r["ai_status"] != "OK"
    ]

    populate_filtered_sheet(
        workbook,
        "Action Required",
        action_records
    )

    populate_filtered_sheet(
        workbook,
        "Deadlines",
        deadline_records
    )

    populate_filtered_sheet(
        workbook,
        "High Priority",
        high_priority_records
    )

    populate_filtered_sheet(
        workbook,
        "Low Confidence",
        low_confidence_records
    )

    # ========================================================
    # ATTACHMENTS
    # ========================================================

    attachment_sheet = workbook.create_sheet("Attachments")

    attachment_sheet.append([
        "Email Date",
        "Sender",
        "Subject",
        "Filename",
        "MIME Type",
        "Size",
        "Age (Days)",
        "AI Category",
        "AI Priority",
        "AI Action",
        "AI Confidence",
    ])

    for record in records:
        if record["attachment_count"] == 0:
            continue

        names = record["attachment_names"].split("; ")
        types = record["attachment_types"].split("; ")

        for index, name in enumerate(names):
            mime_type = (
                types[index]
                if index < len(types)
                else ""
            )

            attachment_sheet.append([
                record["date"],
                record["sender"],
                record["subject"],
                name,
                mime_type,
                record["attachment_size"],
                record["age_days"],
                record["category"],
                record["priority"],
                record["action"],
                record["confidence"],
            ])

    # ========================================================
    # FORMATTING
    # ========================================================

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )
            cell.alignment = Alignment(
                vertical="top"
            )

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        for column in sheet.columns:
            max_length = 0
            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:
                try:
                    length = len(str(cell.value))

                    if length > max_length:
                        max_length = length

                except Exception:
                    pass

            sheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                50
            )

        sheet.auto_filter.ref = sheet.dimensions

    workbook.save(filepath)

    return filepath


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 72)
    print("SRNR EMAIL AGENT - VERSION 3")
    print("EMAIL INTELLIGENCE + THREAD AWARENESS")
    print("=" * 72)

    print(f"\nProject: {PROJECT_DIR}")
    print(f"AI Model: {OLLAMA_MODEL}")
    print(f"Maximum emails this run: {MAX_MESSAGES}")
    print(f"Low-confidence threshold: {LOW_CONFIDENCE_THRESHOLD}")

    print("\nIMPORTANT:")
    print("This version WILL NOT:")
    print("  - Send emails")
    print("  - Delete emails")
    print("  - Move emails")
    print("  - Modify Gmail labels")
    print("  - Archive emails")
    print("  - Download attachments")
    print("  - Copy files to external storage")
    print("  - Take automatic AI actions")

    print(
        "\nIt only reads Gmail and sends email/thread text "
        "to your LOCAL Ollama/Qwen model."
    )

    try:
        # ----------------------------------------------------
        # TEST OLLAMA
        # ----------------------------------------------------

        print("\nChecking Ollama...")

        try:
            ollama_test = requests.get(
                "http://127.0.0.1:11434/api/tags",
                timeout=10
            )

            ollama_test.raise_for_status()

            models = ollama_test.json().get(
                "models",
                []
            )

            installed_models = [
                model.get("name", "")
                for model in models
            ]

            print("Ollama connection: OK")
            print("Installed models:")

            for model in installed_models:
                print(f"  - {model}")

            if OLLAMA_MODEL not in installed_models:
                print("\nWARNING:")
                print(
                    f"Configured model '{OLLAMA_MODEL}' "
                    f"was not found in Ollama."
                )

                print(
                    "Change OLLAMA_MODEL in the script "
                    "to match 'ollama list'."
                )

                return

        except Exception as error:
            print("\nERROR: Cannot connect to Ollama.")
            print("Make sure Ollama is running.")
            print(f"Details: {error}")
            return

        # ----------------------------------------------------
        # GMAIL
        # ----------------------------------------------------

        print("\nConnecting to Gmail...")

        service = authenticate_gmail()

        print("Gmail connection: OK")

        # ----------------------------------------------------
        # MESSAGE IDS
        # ----------------------------------------------------

        message_ids = get_message_ids(service)

        print(
            f"\nEmails selected for processing: "
            f"{len(message_ids):,}"
        )

        if not message_ids:
            print("No emails found.")
            return

        # ----------------------------------------------------
        # PROCESS
        # ----------------------------------------------------

        records = []

        for index, message_id in enumerate(
            message_ids,
            start=1
        ):
            print(
                f"\nProcessing email "
                f"{index}/{len(message_ids)}"
            )

            try:
                record = process_message(
                    service,
                    message_id
                )

                records.append(record)

            except Exception as error:
                print(
                    f"ERROR processing "
                    f"{message_id}: {error}"
                )

        # ----------------------------------------------------
        # CREATE EXCEL
        # ----------------------------------------------------

        if not records:
            print("\nNo emails were successfully processed.")
            return

        print("\nCreating Excel report...")

        report = create_excel_report(records)

        # ----------------------------------------------------
        # COMPLETE
        # ----------------------------------------------------

        print("\n" + "=" * 72)
        print("SRNR EMAIL AGENT V3 - COMPLETE")
        print("=" * 72)

        print(
            f"\nEmails processed: "
            f"{len(records):,}"
        )

        print("\nExcel report:")
        print(report)

        review_count = sum(
            1
            for r in records
            if r["ai_status"] != "OK"
        )

        print(
            f"\nEmails flagged for AI review: "
            f"{review_count:,}"
        )

        print("\nGmail was NOT modified.")
        print("Attachments were NOT downloaded.")
        print("No automatic actions were taken.")

    except Exception as error:
        print("\nERROR:")
        print(error)

        print(
            "\nCheck the error above before continuing."
        )


# ============================================================
# START
# ============================================================

if __name__ == "__main__":
    main()
